
import streamlit as st
import pandas as pd
import numpy as np
from docx import Document
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn
from io import BytesIO
import os
from docx.shared import Pt


def load_substitution_rules(sub_file):
    """Load substitution rules from Excel file (columns: old, new)"""
    try:
        sub_df = pd.read_excel(sub_file, header=None)
        if len(sub_df.columns) >= 2:
            return dict(zip(sub_df[0], sub_df[1]))
        return {}
    except Exception as e:
        st.error(f"Error loading substitution file: {str(e)}")
        return {}


def process_dataframe(df, sub_dict, remove_last_n_rows, remove_cols, round_decimals):
    """Apply all transformations to the dataframe"""
    processed_df = df.copy()

    # Remove rows/columns
    if remove_last_n_rows and remove_last_n_rows > 0:
        processed_df = processed_df.iloc[:-remove_last_n_rows]
    if remove_cols:
        start_col, end_col = remove_cols
        cols_to_drop = processed_df.columns[start_col - 1:end_col]
        processed_df = processed_df.drop(cols_to_drop, axis=1)

    # Apply substitutions to headers first
    if sub_dict:
        processed_df.columns = [str(col) for col in processed_df.columns]
        for old, new in sub_dict.items():
            processed_df.columns = [col.replace(str(old), str(new)) for col in processed_df.columns]

        # Apply substitutions to cell content
        processed_df = processed_df.replace(sub_dict, regex=True)

    # Number formatting
    if round_decimals is not None:
        for col in processed_df.select_dtypes(include=['number']).columns:
            processed_df[col] = processed_df[col].apply(
                lambda x: f"{float(x):,.{round_decimals}f}".replace('.', ',')
                if pd.notna(x) else '-'
            )

    # Always replace NaN with -
    processed_df = processed_df.replace([np.nan, 'nan', 'NaN', 'NaT'], '-')

    return processed_df


def convert_excel_to_word(df):
    """Convert DataFrame to Word document with borders"""
    doc = Document()
    # doc.add_heading('Processed Data', level=1)
    table = doc.add_table(rows=1, cols=len(df.columns))

    def set_font(cell, text):
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
    def set_cell_borders(cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        for border_name in ['top', 'left', 'bottom', 'right']:
            border = OxmlElement(f'w:{border_name}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:space'), '0')
            border.set(qn('w:color'), '000000')
            tcPr.append(border)

    # Add header row
    hdr_cells = table.rows[0].cells
    for i, column in enumerate(df.columns):
        set_font(hdr_cells[i], str(column))
        set_cell_borders(hdr_cells[i])

    # Add data rows
    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        for i, value in enumerate(row):
            set_font(row_cells[i], str(value))
            set_cell_borders(row_cells[i])

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def main():
    st.title("Excel Batch to Word Converter")

    # File uploaders
    col1, col2 = st.columns(2)
    with col1:
        data_files = st.file_uploader("Upload Excel files to convert",
                                    type=["xlsx", "xls"],
                                    accept_multiple_files=True)
    with col2:
        sub_file = st.file_uploader("Upload substitution file (sub.xlsx)",
                                  type=["xlsx", "xls"],
                                  help="First column: text to find, Second column: replacement text")

    if data_files:
        # Load substitution rules
        sub_dict = load_substitution_rules(sub_file) if sub_file else {}

        # Display loaded substitutions
        if sub_dict:
            st.info(f"Loaded {len(sub_dict)} substitution rules")
            if st.checkbox("Show substitution rules"):
                st.dataframe(pd.DataFrame(list(sub_dict.items()), columns=["Find", "Replace"]))

        # Transformation options
        with st.expander("Transformation Options", expanded=True):
            cols = st.columns(3)
            with cols[0]:
                remove_rows = st.checkbox("Remove last N rows")
                if remove_rows:
                    n_rows = st.number_input("Number of rows to remove from end", 1, 100, 1, key="n_rows")
            with cols[1]:
                remove_cols = st.checkbox("Remove columns")
                if remove_cols:
                    col_range = st.slider("Column range to remove", 1, 50, (1, 1), key="col_range")
            with cols[2]:
                round_enabled = st.checkbox("Round numbers")
                round_decimals = st.number_input("Decimal places", 0, 6, 2, key="decimals") if round_enabled else None

        # Process each file
        for data_file in data_files:
            try:
                original_df = pd.read_excel(data_file)
                file_name = os.path.splitext(data_file.name)[0]

                with st.expander(f"Processing: {file_name}", expanded=True):  # Changed to expanded=True
                    # Process data
                    processed_df = process_dataframe(
                        original_df,
                        sub_dict,
                        n_rows if remove_rows else None,
                        col_range if remove_cols else None,
                        round_decimals if round_enabled else None
                    )

                    # Show complete modified preview
                    st.subheader("Complete Modified Table Preview")
                    st.dataframe(processed_df, height=400)  # Increased height to show more rows

                    # Add download button inside the expander
                    word_buffer = convert_excel_to_word(processed_df)
                    st.download_button(
                        label=f"Download {file_name}.docx",
                        data=word_buffer,
                        file_name=f"{file_name}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"download_{file_name}"
                    )

            except Exception as e:
                st.error(f"Error processing {data_file.name}: {str(e)}")


if __name__ == "__main__":
    # Configure Streamlit to show full tables
    st.set_page_config(layout="wide")
    main()